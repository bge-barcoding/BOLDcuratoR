import csv
import io
from pathlib import Path

import pandas as pd
import pytest

from boldcurator.io.annotations import Annotations, merge_annotations
from boldcurator.io import exports


@pytest.fixture
def annotations() -> Annotations:
    a = Annotations()
    a.set_selected("P1", user="curator")
    a.set_flag("P2", "misidentification", user="curator")
    a.set_note("P3", "check the voucher", user="curator")
    return a


def _frame() -> pd.DataFrame:
    return pd.DataFrame([
        {"processid": "P1", "sampleid": "S1", "identification": "Danaus plexippus",
         "species": "Danaus plexippus", "inst": "NHM"},
        {"processid": "P2", "sampleid": "S2", "identification": "Danaus chrysippus",
         "species": "Danaus chrysippus", "inst": "NHM"},
        {"processid": "P3", "sampleid": "S3", "identification": None,
         "species": "Pieris rapae", "inst": ""},
    ])


def test_selection_alone_does_not_count_as_annotated(annotations):
    """R's download_annotated reads flags/updated_ids/notes, not selection.

    Auto-selection would otherwise mark every representative as
    curator-annotated.
    """
    assert annotations.annotated_processids() == {"P2", "P3"}
    assert annotations.selected_processids() == {"P1"}


def test_annotation_columns_lead_the_export(annotations):
    from boldcurator.io.annotations import ANNOTATION_COLUMNS

    merged = merge_annotations(_frame(), annotations)
    assert list(merged.columns[:len(ANNOTATION_COLUMNS)]) == list(ANNOTATION_COLUMNS)
    assert list(merged["selected"]) == [True, False, False]
    assert list(merged["id_flag"]) == ["", "misidentification", ""]
    assert merged.loc[merged["processid"] == "P2", "curator"].iloc[0] == "curator"
    # The snapshot's own identification_method is kept, not replaced.
    assert "identification_method" not in merged.columns
    assert merged.loc[merged["processid"] == "P3", "curator_notes"].iloc[0] == \
        "check the voucher"


def test_institution_column_survives_export(annotations):
    """R whitelists "institution"; the real column is `inst`, so it was dropped."""
    merged = merge_annotations(_frame(), annotations)
    assert "inst" in merged.columns
    assert merged["inst"].iloc[0] == "NHM"


def test_tsv_is_quoted_so_tabs_and_newlines_survive(tmp_path):
    a = Annotations()
    a.set_note("P1", "line one\nline two\tafter a tab")
    merged = merge_annotations(_frame(), a)
    path = exports.write_tsv(merged, tmp_path / "out.tsv", snapshot_id="test")

    # Parse as CSV over the whole stream, not line by line -- a quoted newline
    # is part of one field, so splitlines() would itself break what this is
    # checking.
    text = path.read_text(encoding="utf-8")
    body = "".join(l for l in text.splitlines(keepends=True) if not l.startswith("#"))
    rows = list(csv.reader(io.StringIO(body), delimiter="\t"))
    header, first = rows[0], rows[1]
    assert first[header.index("curator_notes")] == "line one\nline two\tafter a tab"
    assert len(rows) == 4      # header + 3 records, not split by the embedded newline


def test_exports_carry_provenance_and_the_scoring_warning(tmp_path):
    path = exports.write_tsv(_frame(), tmp_path / "p.tsv", snapshot_id="2026-09-11")
    text = path.read_text(encoding="utf-8")
    assert "snapshot: 2026-09-11" in text
    assert "NOT comparable" in text


def test_exports_carry_the_cc_by_sa_attribution(tmp_path):
    """Plan item 0.3: the source data package is CC BY-SA 4.0, which requires

    attribution on redistribution -- an export from this app is exactly that.
    """
    tsv = exports.write_tsv(_frame(), tmp_path / "p.tsv", snapshot_id="x")
    csv_path = exports.write_csv(_frame(), tmp_path / "p.csv", snapshot_id="x")
    for path in (tsv, csv_path):
        text = path.read_text(encoding="utf-8")
        assert "CC BY-SA 4.0" in text
        assert "Barcode of Life Data System" in text


def test_fasta_header_falls_back_through_identification_then_species():
    assert exports.fasta_header(
        pd.Series({"processid": "P1", "identification": "Danaus plexippus"})
    ) == ">P1|Danaus plexippus"
    assert exports.fasta_header(
        pd.Series({"processid": "P2", "identification": None, "species": "Pieris rapae"})
    ) == ">P2|Pieris rapae"
    assert exports.fasta_header(pd.Series({"processid": "P3"})) == ">P3|Unknown"


def test_fasta_writes_streamed_pairs(tmp_path):
    path, n = exports.write_fasta(
        _frame(), [("P1", "ACGT"), ("P2", "TTTT"), ("P3", "")],
        tmp_path / "s.fasta",
    )
    assert n == 2      # the empty sequence is skipped
    assert path.read_text(encoding="utf-8").splitlines() == [
        ">P1|Danaus plexippus", "ACGT", ">P2|Danaus chrysippus", "TTTT"
    ]


def test_curation_report_has_the_narrow_fixed_column_set(annotations):
    report = exports.curation_report(_frame(), annotations)
    assert list(report.columns) == list(exports.CURATION_REPORT_COLUMNS)
    assert set(report["processid"]) == {"P2", "P3"}


def test_export_all_writes_everything_and_says_what_it_skipped(store, tmp_path):
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    a = Annotations()
    first = str(result.specimens["processid"].iloc[0])
    a.set_flag(first, contamination=True, sequence_taxon="Homo sapiens",
               user="tester")

    out = exports.export_all(result, tmp_path, annotations=a, store=store)
    assert {"all", "annotated", "curation_report", "search_results",
            "bin_analysis", "fasta"} <= set(out.written)
    assert "selected" in out.skipped
    for path in out.written.values():
        assert Path(path).exists() and Path(path).stat().st_size > 0


def test_export_all_skips_fasta_without_a_snapshot(store, tmp_path):
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    out = exports.export_all(result, tmp_path, store=None)
    assert "fasta" in out.skipped
    assert "sequences cannot be fetched" in out.skipped["fasta"]


def test_bin_analysis_workbook_has_three_populated_sheets(store, tmp_path):
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    path = exports.write_bin_analysis_xlsx(
        result.bin_analysis, tmp_path / "bins.xlsx", snapshot_id="x"
    )
    sheets = pd.read_excel(path, sheet_name=None)
    assert set(sheets) == {"Summary", "Content", "Statistics"}
    # R's version leaves Summary and Statistics empty; both carry data here.
    assert all(len(frame) > 0 for frame in sheets.values())
    summary_text = sheets["Summary"].to_string()
    assert "CC BY-SA 4.0" in summary_text


def test_species_analysis_workbook_has_the_checklist_and_gap_analysis(store, tmp_path):
    """Plan round 3, items 3-4: one xlsx download for the Species screen."""
    from boldcurator.core.pipeline import run_search
    from boldcurator.core.summaries import build_species_checklist, gap_analysis

    result = run_search(store, taxa_text="Danaus plexippus")
    checklist = build_species_checklist(result.specimens, result.bags_grades)
    gaps = gap_analysis(result.taxonomy_groups, result.specimens)

    path = exports.write_species_analysis_xlsx(
        checklist, gaps, tmp_path / "species.xlsx", snapshot_id="x")
    sheets = pd.read_excel(path, sheet_name=None)
    assert set(sheets) == {"Summary", "Species checklist", "Gap analysis"}
    assert len(sheets["Species checklist"]) == len(checklist)
    assert len(sheets["Gap analysis"]) == len(gaps)
    assert "mean_quality_score" not in sheets["Species checklist"].columns
    summary_text = sheets["Summary"].to_string()
    assert "CC BY-SA 4.0" in summary_text


def test_species_analysis_workbook_survives_an_empty_gap_analysis(store, tmp_path):
    """A dataset/project-code-only search has no taxa typed -- gaps is
    legitimately empty, and that must not be treated as nothing to export."""
    from boldcurator.core.summaries import build_species_checklist

    checklist = build_species_checklist(
        pd.DataFrame({"species": ["Danaus plexippus"], "bin_uri": ["BOLD:AAA0001"],
                     "country.ocean": ["Canada"], "quality_score": [10]}),
        pd.DataFrame(columns=["species", "bags_grade"]),
    )
    gaps = pd.DataFrame(columns=["input_taxon", "status", "matched_species",
                                 "specimen_count", "notes"])

    path = exports.write_species_analysis_xlsx(
        checklist, gaps, tmp_path / "species.xlsx", snapshot_id="x")
    sheets = pd.read_excel(path, sheet_name=None)
    assert len(sheets["Species checklist"]) == 1
    assert len(sheets["Gap analysis"]) == 0


def test_search_state_exports_species_analysis_end_to_end(store, tmp_path):
    from boldcurator.ui.state import AppState

    state = AppState(store, page_size=25)
    state.run_search("Danaus plexippus")
    written = state.search.export_species_analysis(store, tmp_path / "s.xlsx")
    assert written is not None
    sheets = pd.read_excel(written, sheet_name=None)
    assert set(sheets) == {"Summary", "Species checklist", "Gap analysis"}
    assert len(sheets["Species checklist"]) > 0
    assert "Danaus plexippus" in sheets["Gap analysis"]["input_taxon"].tolist()


def test_unknown_flag_is_refused():
    a = Annotations()
    with pytest.raises(ValueError, match="Unknown record flag"):
        a.set_flag("P1", "not_a_real_flag")
    # Name issue and synonym belong to the name, not one record.
    with pytest.raises(ValueError, match="taxon updates"):
        a.set_flag("P1", "synonym")


def test_auto_selections_reach_the_selected_export(store, tmp_path):
    """R writes auto-selections into the same store download_selected reads."""
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    assert result.selections, "pipeline should have auto-selected representatives"

    out = exports.export_all(result, tmp_path, store=store)
    assert "selected" in out.written
    assert "selected_fasta" in out.written
    assert "selected" not in out.skipped


def test_explicit_annotations_override_auto_selection(store, tmp_path):
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    out = exports.export_all(result, tmp_path, annotations=Annotations(), store=store)
    assert "selected" in out.skipped


def test_xlsx_export_holds_the_rows_and_the_provenance(tmp_path, annotations):
    """Round 9, item 4: the Specimens tab's All / Selected / Annotated
    downloads are workbooks -- the rows first, the provenance and licence
    the TSV carries in comments on an About sheet."""
    from openpyxl import load_workbook

    merged = merge_annotations(_frame(), annotations)
    path = exports.write_xlsx(merged, tmp_path / "out.xlsx", snapshot_id="2026-09-11")
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == ["Records", "About"]
    rows = list(book["Records"].values)
    assert list(rows[0]) == list(merged.columns)
    assert len(rows) == 4
    pids = [r[list(merged.columns).index("processid")] for r in rows[1:]]
    assert pids == ["P1", "P2", "P3"]
    about = " ".join(str(r[0]) for r in book["About"].values)
    assert "snapshot: 2026-09-11" in about
    assert "CC BY-SA 4.0" in about


def test_xlsx_export_never_writes_a_value_as_a_formula(tmp_path):
    """Snapshot text and curator notes are data: one starting with "=" must
    reach the spreadsheet as text, not as a formula Excel would run."""
    from openpyxl import load_workbook

    frame = pd.DataFrame({"processid": ["P1", "P2"],
                          "notes": ['=HYPERLINK("http://x","y")', None],
                          "count": [3, pd.NA],
                          "bad": ["tab\tok, bell\x07 gone", "fine"]})
    path = exports.write_xlsx(frame, tmp_path / "f.xlsx")
    sheet = load_workbook(path)["Records"]
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("http://x","y")'
    assert sheet["B3"].value is None
    assert sheet["C2"].value == 3 and sheet["C3"].value is None
    assert sheet["D2"].value == "tab\tok, bell gone"


def test_processid_list_is_one_per_line_with_no_header(tmp_path):
    """Round 9, item 3: pasted straight into a BOLD dataset."""
    path = exports.write_processid_list(["P1", "P2"], tmp_path / "ids.txt")
    assert path.read_text(encoding="utf-8") == "P1\nP2\n"


def test_xlsx_header_is_bold_and_frozen_and_progress_is_reported(tmp_path):
    """Round 10, item 3: the hand-written workbook (openpyxl was ten times
    slower) is still a proper one -- bold, frozen header -- and says how far
    it has got, for the download's progress line."""
    import numpy as np
    from openpyxl import load_workbook

    n = exports._XLSX_CHUNK_ROWS * 2 + 5
    frame = pd.DataFrame({"processid": [f"P{i}" for i in range(n)],
                          "count": np.arange(n, dtype=np.int64),
                          "score": np.linspace(0, 1, n),
                          "flag": [np.bool_(i % 2) for i in range(n)],
                          "big": [float("inf")] + [1.5] * (n - 1),
                          "long": ["x" * 40_000] + [""] * (n - 1)})
    seen = []
    path = exports.write_xlsx(frame, tmp_path / "p.xlsx",
                              progress=lambda done, total: seen.append((done, total)))
    assert seen[-1] == (n, n) and len(seen) == 3
    sheet = load_workbook(path)["Records"]
    assert sheet.freeze_panes == "A2" and sheet["A1"].font.b
    assert (sheet["B3"].value, sheet["D2"].value, sheet["D3"].value) == (1, False, True)
    assert sheet["E2"].value is None and sheet["E3"].value == 1.5    # inf: empty
    assert len(sheet["F2"].value) == 32_767 and sheet["F3"].value is None
    assert sheet.max_row == n + 1


def test_an_annotations_snapshot_does_not_follow_later_changes(annotations):
    """Round 10, item 3: a download prepared in the background writes the
    annotations as they were when it was clicked."""
    copy = annotations.snapshot()
    annotations.set_flag("P1", "id_unfounded")
    annotations.set_taxon_update("Aus bis", "synonym", "Aus bus")
    annotations.set_selected("P9")
    annotations.set_working("P3")
    assert "P1" not in copy.flags and "P9" not in copy.selected
    assert not copy.taxon_updates
    assert not copy.working
    assert copy.to_dict()["curator_notes"] == {"P3": annotations.curator_notes["P3"]}


def _taxon_annotations() -> Annotations:
    a = Annotations()
    a.set_taxon_update("Danaus plexippus", "synonym", "Danaus archippus",
                       rank="species", citation="Smith 2019", user="Ann Curator")
    return a


def test_records_a_taxon_update_reaches_are_in_the_curation_report():
    """Agreed: one row per affected record, its updated ID and where it came
    from -- and a taxon updates table beside it."""
    a = _taxon_annotations()
    a.set_flag("P2", "misidentification", user="Ann Curator")
    a.set_updated_id("P2", "Danaus plexippus", rank="species",
                     method="Morphology", user="Ann Curator")
    report = exports.curation_report(_frame(), a)
    rows = report.set_index("processid")
    assert set(rows.index) == {"P1", "P2"}
    assert rows.loc["P1", "updated_id"] == "Danaus archippus"
    assert rows.loc["P1", "update_source"] == "taxon (synonym)"
    assert rows.loc["P1", "updated_id_method"] == ""
    # P2's own misidentification is to the junior name, so the synonym
    # follows it -- and keeps the record's own method.
    assert rows.loc["P2", "updated_id"] == "Danaus archippus"
    assert rows.loc["P2", "update_source"] == "record → taxon (synonym)"
    assert rows.loc["P2", "id_flag"] == "misidentification"
    assert rows.loc["P2", "updated_id_method"] == "Morphology"

    taxa = exports.taxon_updates_report(_frame(), a)
    assert taxa.to_dict("records") == [{
        "old_name": "Danaus plexippus", "flag": "synonym",
        "new_name": "Danaus archippus", "rank": "species",
        "citation": "Smith 2019", "note": "", "curator": "Ann Curator",
        "timestamp": a.taxon_updates["Danaus plexippus"]["timestamp"],
        "records_affected": 1}]


def test_contamination_alone_reports_the_voucher_id_unchanged():
    """The doc: the voucher's correct ID goes in the updated identification
    field -- with no misidentification, that is its current ID."""
    a = Annotations()
    a.set_flag("P1", contamination=True, sequence_taxon="Homo sapiens")
    report = exports.curation_report(_frame(), a).set_index("processid")
    assert report.loc["P1", "updated_id"] == "Danaus plexippus"
    assert report.loc["P1", "sequence_taxon"] == "Homo sapiens"
    assert bool(report.loc["P1", "contamination"]) is True


def test_export_all_writes_the_taxon_updates(store, tmp_path):
    from boldcurator.core.pipeline import run_search

    result = run_search(store, taxa_text="Danaus plexippus")
    out = exports.export_all(result, tmp_path, annotations=_taxon_annotations(),
                             store=store)
    assert "taxon_updates" in out.written and "annotated" in out.written
    text = Path(out.written["taxon_updates"]).read_text(encoding="utf-8")
    assert "Danaus archippus" in text
    assert "taxon_updates" in exports.export_all(
        result, tmp_path / "none", store=store).skipped


def test_the_annotated_workbook_carries_a_taxon_updates_sheet(store, tmp_path):
    from boldcurator.ui.state import AppState

    state = AppState(store, page_size=25)
    state.run_search("Danaus plexippus")
    state.annotations.set_taxon_update("Danaus plexippus", "name_issue",
                                       "Danaus plexipus", rank="species")
    written = state.search.export_specimens(store, "annotated",
                                            tmp_path / "a.xlsx", excel=True)
    sheets = pd.read_excel(written, sheet_name=None)
    assert list(sheets) == ["Records", "Taxon updates", "About"]
    assert sheets["Taxon updates"]["new_name"].tolist() == ["Danaus plexipus"]
    assert (sheets["Records"]["updated_id"] == "Danaus plexipus").all()


def test_downloads_carry_the_chosen_views_grade_beside_the_original(store, tmp_path):
    """Round 12, item 3: before -> after in the downloads too."""
    from boldcurator.ui.state import AppState

    state = AppState(store, page_size=25)
    state.run_search("Danaus plexippus")
    frame = state.search.analysis(store).specimens
    pid = str(frame["processid"].iloc[0])
    state.annotations.set_flag(pid, contamination=True, sequence_taxon="Homo sapiens")
    written = state.search.export_specimens(store, "annotated", tmp_path / "a.tsv")
    assert "bags_grade_unaltered" not in Path(written).read_text(encoding="utf-8")
    state.search.grade_view = "unaltered"
    written = state.search.export_specimens(store, "annotated", tmp_path / "b.tsv")
    body = [l for l in Path(written).read_text(encoding="utf-8").splitlines()
            if not l.startswith("#")]
    header = body[0].split("\t")
    assert header.index("bags_grade_unaltered") == header.index("bags_grade") + 1
    # The contaminated record is left out of that view, so it has no grade there.
    assert body[1].split("\t")[header.index("bags_grade_unaltered")] in ("", '""')
